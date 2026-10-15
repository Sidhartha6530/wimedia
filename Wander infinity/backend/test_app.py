import json
import os
import tempfile
import unittest
from unittest.mock import patch

from app import app
from openpyxl import load_workbook


class ContactStoreTests(unittest.TestCase):
    def setUp(self):
        app.config['TESTING'] = True
        self.client = app.test_client()
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.storage_path = os.path.join(self.temp_dir.name, 'contacts.json')
        data_file_patcher = patch('app.DATA_FILE', self.storage_path)
        data_file_patcher.start()
        self.addCleanup(data_file_patcher.stop)

    def test_contact_is_saved_to_backend_file(self):
        payload = {
            'firstName': 'Alice',
            'lastName': 'Smith',
            'email': 'alice@example.com',
            'phone': '+1 555 0100',
            'company': 'Acme Studio',
            'website': '=HYPERLINK("https://example.com","Click")',
            'projectType': 'Brand strategy',
            'budget': '₹50,000 – ₹2,00,000',
            'timeline': 'Within 1 month',
            'preferredContact': 'Email',
            'location': 'Hyderabad, IST',
            'referralSource': 'Referral',
            'subject': 'Project inquiry',
            'message': 'I want to build a website.'
        }

        response = self.client.post('/api/contact', json=payload)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(os.path.exists(self.storage_path))

        with open(self.storage_path, 'r', encoding='utf-8') as fh:
            saved = json.load(fh)

        self.assertEqual(len(saved), 1)
        self.assertEqual(saved[0]['email'], 'alice@example.com')
        self.assertEqual(saved[0]['phone'], '+1 555 0100')
        self.assertEqual(saved[0]['company'], 'Acme Studio')
        self.assertEqual(saved[0]['website'], '=HYPERLINK("https://example.com","Click")')
        self.assertEqual(saved[0]['projectType'], 'Brand strategy')
        self.assertEqual(saved[0]['budget'], '₹50,000 – ₹2,00,000')
        self.assertEqual(saved[0]['timeline'], 'Within 1 month')
        self.assertEqual(saved[0]['preferredContact'], 'Email')
        self.assertEqual(saved[0]['location'], 'Hyderabad, IST')
        self.assertEqual(saved[0]['referralSource'], 'Referral')
        self.assertEqual(saved[0]['message'], 'I want to build a website.')

        excel_path = os.path.splitext(self.storage_path)[0] + '.xlsx'
        workbook = load_workbook(excel_path, read_only=True, data_only=False)
        self.addCleanup(workbook.close)
        worksheet = workbook['Contacts']
        self.assertEqual(worksheet['A1'].value, 'name')
        self.assertEqual(worksheet['B2'].value, 'alice@example.com')
        self.assertEqual(worksheet['E2'].value, '=HYPERLINK("https://example.com","Click")')
        self.assertEqual(worksheet['E2'].data_type, 's')
        self.assertEqual(worksheet['M2'].value, 'I want to build a website.')
        self.assertEqual(worksheet['M2'].data_type, 's')

    def test_existing_contacts_are_exported_when_listed(self):
        existing_contacts = [{
            'name': 'Existing Contact',
            'email': 'existing@example.com',
            'message': 'Previously saved',
        }]
        with open(self.storage_path, 'w', encoding='utf-8') as fh:
            json.dump(existing_contacts, fh)

        response = self.client.get('/api/contacts')

        self.assertEqual(response.status_code, 200)
        excel_path = os.path.splitext(self.storage_path)[0] + '.xlsx'
        self.assertTrue(os.path.exists(excel_path))
        workbook = load_workbook(excel_path, read_only=True)
        self.addCleanup(workbook.close)
        worksheet = workbook['Contacts']
        self.assertEqual(worksheet['A2'].value, 'Existing Contact')
        self.assertEqual(worksheet['B2'].value, 'existing@example.com')
        self.assertEqual(worksheet['M2'].value, 'Previously saved')


if __name__ == '__main__':
    unittest.main()
